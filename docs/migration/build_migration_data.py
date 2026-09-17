# -*- coding: utf-8 -*-
"""
現行Excel「（仮）NKC棚卸.xlsx」シート1 → Web移行データ(CSV)生成スクリプト
- 確定ルール(docs/03_移行ルール.md R-01〜R-10、未確定-1〜6 の回答)に従う
- 出力先: docs/migration/*.csv (UTF-8 BOM付き)
- 実行: python3 docs/migration/build_migration_data.py
"""
import csv, os, sys, datetime, collections
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(BASE, '..', 'source', '（仮）NKC棚卸.xlsx')
SHEET = 'シート1'
HEADER_ROW = 3
FIRST_ROW, LAST_ROW = 4, 114
IMPORT_USER = 'excel_import'
NOW = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')

# 個体管理にする行(未確定-4: 管理Noのある6行のみ)
UNIT_ROWS = {4, 5, 34, 35, 64, 67}

wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]

# --- ヘッダー確認(列位置がずれていたら停止) ---
expected = {2: '品名', 3: '状態（新品、中古）', 4: '管理No', 5: '保管場所', 6: '棚卸数', 7: '棚卸数 2026/1/6', 8: '備考'}
for col, name in expected.items():
    v = ws.cell(HEADER_ROW, col).value
    if v != name:
        sys.exit(f'ヘッダー不一致: {ws.cell(HEADER_ROW, col).coordinate}={v!r} expected {name!r}')

# --- 結合セル展開(R-04): 保管場所(E列)は各行に値を持たせる。G72:G73 は未確定-1 (a) → 先頭行のみ ---
merged_loc = {}
for rng in ws.merged_cells.ranges:
    if rng.min_col == rng.max_col == 5:
        v = ws.cell(rng.min_row, 5).value
        for r in range(rng.min_row, rng.max_row + 1):
            merged_loc[r] = v

def s(v):
    """原文保持(R-02): 文字列はそのまま、数値は整数表記の文字列、Noneは''"""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)

def qty(v):
    """棚卸数: 空欄は空欄のまま(R-05)。数値は整数"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v)
    sys.exit(f'棚卸数に数値以外の値: {v!r}')

# --- Excel行の読み込み ---
rows = []
for r in range(FIRST_ROW, LAST_ROW + 1):
    name = ws.cell(r, 2).value
    cond = ws.cell(r, 3).value
    mgmt = ws.cell(r, 4).value
    loc = merged_loc.get(r, ws.cell(r, 5).value)
    f = ws.cell(r, 6).value
    g = ws.cell(r, 7).value            # 結合G72:G73 は G72 のみ値あり(未確定-1 (a))
    note = ws.cell(r, 8).value
    col_i = ws.cell(r, 9).value
    if all(v is None for v in (name, cond, mgmt, loc, f, g, note, col_i)):
        continue
    # R-01: I列は備考へ追記(H列があれば改行で連結)
    notes = s(note)
    if col_i is not None:
        notes = (notes + '\n' + s(col_i)) if notes else s(col_i)
    rows.append(dict(row=r, name=s(name), cond=s(cond), mgmt=s(mgmt), loc=s(loc),
                     f=qty(f), g=qty(g), notes=notes))
assert len(rows) == 111, len(rows)

# --- マスタ ---
conditions = [('new', '新品', 1), ('used', '中古', 2)]
cond_code = {n: c for c, n, _ in conditions}

loc_names = []
for r in rows:
    if r['loc'] and r['loc'] not in loc_names:
        loc_names.append(r['loc'])
locations = [(i + 1, n, i + 1) for i, n in enumerate(loc_names)]
loc_id = {n: i for i, n, _ in locations}

categories = ['PC本体', 'ディスプレイ', '入力機器', '周辺機器', 'ストレージ・メディア', 'ネットワーク機器',
              'ケーブル・変換アダプタ', '電源関連', '工具・測定器', '配線資材・消耗品', '設置・固定器具', 'その他']

# --- 品目(B案): 個体管理行は「品名+状態+保管場所」が同じなら1品目にまとめ、個体を別テーブルへ ---
items = []          # dict
units = []
item_key_to_id = {}
row_to_item = {}
for r in rows:
    is_unit = r['row'] in UNIT_ROWS
    if is_unit:
        key = ('unit', r['name'], r['cond'], r['loc'])
    else:
        key = ('row', r['row'])
    if key in item_key_to_id:
        iid = item_key_to_id[key]
        it = items[iid - 1]
        it['source_excel_rows'].append(r['row'])
        # 備考はまとめる(原文保持: 空でなければ改行連結)
        if r['notes']:
            it['notes'] = (it['notes'] + '\n' + r['notes']) if it['notes'] else r['notes']
    else:
        iid = len(items) + 1
        item_key_to_id[key] = iid
        items.append(dict(
            id=iid, item_code=f'ITEM-{iid:06d}', item_name=r['name'],
            management_type='unit' if is_unit else 'quantity',
            category_id=None,                          # 未確定-5: 移行時は設定しない
            condition_code=cond_code.get(r['cond']) if r['cond'] else None,
            condition_raw=r['cond'],                   # 原文(新品/中古以外が来た場合の保険。今回は新品/中古/空欄のみ)
            stock_type='internal', customer_id=None, manufacturer=None, model_number=None,
            serial_number=None, network_info=None,
            location_id=loc_id.get(r['loc']) if r['loc'] else None,
            notes=r['notes'], is_active=1, sort_order=iid,
            source_excel_rows=[r['row']],
        ))
    row_to_item[r['row']] = iid
    if is_unit:
        units.append(dict(id=len(units) + 1, item_id=iid, management_no=r['mgmt'], serial_number=None,
                          ip_address=None, status='in_stock',
                          location_id=loc_id.get(r['loc']) if r['loc'] else None,
                          notes='', is_active=1, source_excel_row=r['row']))

# 状態が新品/中古以外の値だった場合は停止(今回は該当なし)
for it in items:
    if it['condition_raw'] and it['condition_code'] is None:
        sys.exit(f"状態に未知の値: {it['condition_raw']!r} (rows {it['source_excel_rows']})")

# --- 棚卸(2026-09-17 指示: 棚卸1は棚卸期間・入力日付とも空欄。棚卸2は棚卸期間=2025/12/26(H1 の日付部分)、入力日付=2026/1/8(H2)) ---
h2 = ws['H2'].value
counts = [
    dict(id=1, count_name='2025/12/26 棚卸', base_date='2025-12-26', count_period=None, entry_date=None,
         start_date=None, end_date=None, status='confirmed', notes='Excel シート1 F列「棚卸数」より移行'),
    dict(id=2, count_name='2026/1/6 棚卸', base_date='2026-01-06', count_period='2025/12/26',
         entry_date=h2.strftime('%Y-%m-%d') if isinstance(h2, datetime.datetime) else s(h2),
         start_date=None, end_date=None, status='confirmed', notes='Excel シート1 G列「棚卸数 2026/1/6」より移行'),
]
assert counts[1]['entry_date'] == '2026-01-08', counts[1]['entry_date']

# --- 棚卸明細(R-05: 空欄は明細を作らない) ---
details = []
unit_results = []
for cnt, key in ((1, 'f'), (2, 'g')):
    per_item = collections.OrderedDict()
    for r in rows:
        q = r[key]
        if q is None:
            continue
        iid = row_to_item[r['row']]
        per_item.setdefault(iid, {'qty': 0, 'rows': []})
        per_item[iid]['qty'] += q
        per_item[iid]['rows'].append(r['row'])
        if r['row'] in UNIT_ROWS:
            uid = next(u['id'] for u in units if u['source_excel_row'] == r['row'])
            unit_results.append(dict(id=len(unit_results) + 1, count_id=cnt, unit_id=uid,
                                     result='present' if q >= 1 else 'absent', notes='', source_excel_row=r['row']))
    for iid, d in per_item.items():
        it = items[iid - 1]
        details.append(dict(id=len(details) + 1, count_id=cnt, item_id=iid, count_quantity=d['qty'],
                            confirm_status='confirmed', location_id_at_count=it['location_id'],
                            counted_by=None, counted_at=None, notes='',
                            source_excel_rows=d['rows']))

# --- 出力 ---
def write(name, header, recs):
    p = os.path.join(BASE, name)
    with open(p, 'w', newline='', encoding='utf-8-sig') as fp:
        w = csv.writer(fp)
        w.writerow(header)
        for rec in recs:
            w.writerow(['' if v is None else (';'.join(map(str, v)) if isinstance(v, list) else v) for v in rec])
    print(f'{name}: {len(recs)} rows')

write('conditions.csv', ['code', 'name', 'sort_order', 'is_active'], [(c, n, o, 1) for c, n, o in conditions])
write('locations.csv', ['id', 'name', 'sort_order', 'is_active'], [(i, n, o, 1) for i, n, o in locations])
write('categories.csv', ['id', 'name', 'sort_order', 'is_active'], [(i + 1, n, i + 1, 1) for i, n in enumerate(categories)])
write('customers.csv', ['id', 'name', 'is_active'], [])   # R-06: 移行時は登録しない(候補: 恵会/永和会/腎愛会/エミフル)
write('inventory_items.csv',
      ['id', 'item_code', 'item_name', 'management_type', 'category_id', 'condition_code', 'stock_type', 'customer_id',
       'manufacturer', 'model_number', 'serial_number', 'network_info', 'location_id', 'notes', 'is_active', 'sort_order',
       'created_at', 'updated_at', 'created_by', 'updated_by', 'source_excel_rows'],
      [(it['id'], it['item_code'], it['item_name'], it['management_type'], it['category_id'], it['condition_code'],
        it['stock_type'], it['customer_id'], it['manufacturer'], it['model_number'], it['serial_number'], it['network_info'],
        it['location_id'], it['notes'], it['is_active'], it['sort_order'], NOW, NOW, IMPORT_USER, IMPORT_USER,
        it['source_excel_rows']) for it in items])
write('inventory_units.csv',
      ['id', 'item_id', 'management_no', 'serial_number', 'ip_address', 'status', 'location_id', 'notes', 'is_active',
       'created_at', 'updated_at', 'created_by', 'updated_by', 'source_excel_row'],
      [(u['id'], u['item_id'], u['management_no'], u['serial_number'], u['ip_address'], u['status'], u['location_id'],
        u['notes'], u['is_active'], NOW, NOW, IMPORT_USER, IMPORT_USER, u['source_excel_row']) for u in units])
write('inventory_counts.csv',
      ['id', 'count_name', 'base_date', 'count_period', 'entry_date', 'start_date', 'end_date', 'status', 'notes',
       'created_by', 'created_at', 'confirmed_by', 'confirmed_at'],
      [(c['id'], c['count_name'], c['base_date'], c['count_period'], c['entry_date'], c['start_date'], c['end_date'],
        c['status'], c['notes'], IMPORT_USER, NOW, IMPORT_USER, NOW) for c in counts])
write('inventory_count_details.csv',
      ['id', 'count_id', 'item_id', 'count_quantity', 'confirm_status', 'location_id_at_count', 'counted_by', 'counted_at',
       'notes', 'created_at', 'updated_at', 'created_by', 'updated_by', 'source_excel_rows'],
      [(d['id'], d['count_id'], d['item_id'], d['count_quantity'], d['confirm_status'], d['location_id_at_count'],
        d['counted_by'], d['counted_at'], d['notes'], NOW, NOW, IMPORT_USER, IMPORT_USER, d['source_excel_rows']) for d in details])
write('inventory_count_unit_results.csv',
      ['id', 'count_id', 'unit_id', 'result', 'notes', 'created_at', 'updated_at', 'created_by', 'updated_by', 'source_excel_row'],
      [(u['id'], u['count_id'], u['unit_id'], u['result'], u['notes'], NOW, NOW, IMPORT_USER, IMPORT_USER, u['source_excel_row'])
       for u in unit_results])

# --- 突合(Excel原本 vs 移行データ) ---
print('--- 突合 ---')
ex_f = sum(r['f'] for r in rows if r['f'] is not None)
ex_g = sum(r['g'] for r in rows if r['g'] is not None)
mg_f = sum(d['count_quantity'] for d in details if d['count_id'] == 1)
mg_g = sum(d['count_quantity'] for d in details if d['count_id'] == 2)
print(f'Excel行数 {len(rows)} → 品目 {len(items)} + 個体 {len(units)}(品目にまとめた行: {len(rows) - len(items)})')
print(f'12/26 数量合計 Excel={ex_f} 移行={mg_f} 明細数={sum(1 for d in details if d["count_id"] == 1)} (Excel数値セル数={sum(1 for r in rows if r["f"] is not None)})')
print(f'1/6   数量合計 Excel={ex_g} 移行={mg_g} 明細数={sum(1 for d in details if d["count_id"] == 2)} (Excel数値セル数={sum(1 for r in rows if r["g"] is not None)})')
print(f'備考あり Excel(H or I)={sum(1 for r in rows if r["notes"])} 品目notesあり={sum(1 for it in items if it["notes"])}')
print(f'管理No Excel={sorted(r["mgmt"] for r in rows if r["mgmt"])} 個体={sorted(u["management_no"] for u in units)}')
assert ex_f == mg_f and ex_g == mg_g
covered = sorted(x for it in items for x in it['source_excel_rows'])
assert covered == [r['row'] for r in rows], '行の欠落'
print('OK: 全111行が品目に対応、数量合計一致')
